"""
ICA Geoconsultores - Sincronizador Local M365
Sincroniza los registros generados en la base de datos de SharePoint / Timesheet_App
con la hoja oficial 'Registro_Horas_Detalle' de Timesheet_Maestro_ICA_Estandarizado.xlsx
"""

import os
import json
import openpyxl
from datetime import datetime

SHAREPOINT_DIR = r"C:\Users\ealva\SEID\ICAGEO CHILE - Pegas\Timesheet"
DB_DIR = os.path.join(SHAREPOINT_DIR, "BD_Timesheet_SharePoint")
JSON_DB_PATH = os.path.join(DB_DIR, "timesheet_records.json")
EXCEL_MASTER_PATH = os.path.join(SHAREPOINT_DIR, "Timesheet_Maestro_ICA_Estandarizado.xlsx")

def sync_records():
    if not os.path.exists(JSON_DB_PATH):
        print(f"Error: No se encontró {JSON_DB_PATH}")
        return

    if not os.path.exists(EXCEL_MASTER_PATH):
        print(f"Error: No se encontró el libro maestro {EXCEL_MASTER_PATH}")
        return

    with open(JSON_DB_PATH, "r", encoding="utf-8") as f:
        records = json.load(f)

    wb = openpyxl.load_workbook(EXCEL_MASTER_PATH)
    if "Registro_Horas_Detalle" not in wb.sheetnames:
        print("Error: Hoja 'Registro_Horas_Detalle' no encontrada en el libro maestro.")
        return

    sheet = wb["Registro_Horas_Detalle"]

    # Identificar IDs ya existentes en la hoja
    existing_ids = set()
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if row and row[0]:
            existing_ids.add(str(row[0]).strip())

    inserted_count = 0
    next_row = sheet.max_row + 1

    for rec in records:
        rec_id = rec.get("id")
        if rec_id in existing_ids:
            continue

        fecha_str = rec.get("fecha", "")
        # Determinar día de la semana
        try:
            d_obj = datetime.strptime(fecha_str, "%Y-%m-%d")
            dias = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
            dia_semana = dias[d_obj.weekday()]
        except Exception:
            dia_semana = "Lunes"

        colaborador = rec.get("usuario_nombre", "")
        correo = rec.get("usuario_correo", "")

        for tarea in rec.get("tareas", []):
            proyecto = tarea.get("proyecto", "")
            horas = tarea.get("horas", 0)
            detalle = tarea.get("detalle", "") or tarea.get("actividad", "")
            
            # Fila: [ID_Registro, Fecha_Semana, Fecha_Labor, Consultor, Correo_Corporativo, Proyecto, Codigo_OT, Actividad_Descripcion, Dia_Semana, Horas_HH]
            sheet.cell(row=next_row, column=1, value=rec_id)
            sheet.cell(row=next_row, column=2, value=fecha_str)
            sheet.cell(row=next_row, column=3, value=fecha_str)
            sheet.cell(row=next_row, column=4, value=colaborador)
            sheet.cell(row=next_row, column=5, value=correo)
            sheet.cell(row=next_row, column=6, value=proyecto)
            sheet.cell(row=next_row, column=7, value="OT-AUTO")
            sheet.cell(row=next_row, column=8, value=detalle)
            sheet.cell(row=next_row, column=9, value=dia_semana)
            sheet.cell(row=next_row, column=10, value=horas)
            
            next_row += 1
            inserted_count += 1
            existing_ids.add(rec_id)

    wb.save(EXCEL_MASTER_PATH)
    print(f"Sincronización exitosa: {inserted_count} filas insertadas en {EXCEL_MASTER_PATH}")

if __name__ == "__main__":
    sync_records()
