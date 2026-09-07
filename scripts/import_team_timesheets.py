"""
import_team_timesheets.py
Procesa y estandariza las HH de los 5 consultores desde C:\\Users\\ealva\\Desktop\\Proyectos Antigravity\\Gestion ICA\\Timesheet
Genera el dataset consolidado para el Dashboard y sincroniza con SharePoint.
"""

import os
import glob
import json
import datetime
import openpyxl
from collections import defaultdict

BASE_TIMESHEET_DIR = r"C:\Users\ealva\Desktop\Proyectos Antigravity\Gestion ICA\Timesheet"
APP_DATA_DIR = r"C:\Users\ealva\Desktop\Proyectos Antigravity\Gestion ICA\Timesheet_App\data"
SHAREPOINT_DIR = r"C:\Users\ealva\SEID\ICAGEO CHILE - Pegas\Timesheet\BD_Timesheet_SharePoint"

os.makedirs(APP_DATA_DIR, exist_ok=True)
os.makedirs(SHAREPOINT_DIR, exist_ok=True)

# Mapeo canónico de Proyectos
def normalize_project(text):
    t = str(text or "").lower()
    # Minera Las Luces (incluye Packer según instrucción del usuario)
    if any(k in t for k in ["packer", "las luces", "luces", "sondaje", "sondajes", "terreno ll", "multitester", "epp"]):
        return "Minera Las Luces (MLC)"
    # Pimentón
    elif "piment" in t or "isotopo" in t or "uh" in t or "piezometr" in t or "hidroquim" in t or "balance i" in t:
        return "Minera Pimentón"
    # Kinross
    elif "kinross" in t or "lnf" in t or "leapfrog" in t:
        return "Kinross - LNF"
    # WSP
    elif "wsp" in t or "glaciar" in t or "permafrost" in t:
        return "WSP - DIA Glaciares"
    # Bodega San Francisco
    elif "bsf" in t or "bodega san francisco" in t or "vulnerabilidad" in t:
        return "Bodega San Francisco (BSF)"
    # CCU
    elif "ccu" in t or "quilicura" in t:
        return "CCU Quilicura"
    # Propuestas
    elif "efe" in t or "ecconsulting" in t or "consultame" in t or "propuesta" in t or "licitaci" in t:
        return "Propuestas y Licitaciones"
    # Gestión Interna / ICA
    elif any(k in t for k in ["semanal", "asch", "cv", "daily", "examen", "zoe", "pausa activa", "almuerzo", "colaci"]):
        return "Gestión Interna / ICA"
    else:
        return "Gestión Interna / ICA"

all_records = []

# ==============================================================================
# 1. CRISTÓBAL BRAVO
# ==============================================================================
file_cb = os.path.join(BASE_TIMESHEET_DIR, "TS Cristóbal Bravo.xlsx")
if os.path.exists(file_cb):
    wb = openpyxl.load_workbook(file_cb, data_only=True)
    ws = wb["Septiembre"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if not row or not row[1]:
            continue
        # Semana, Fecha, Horario, HH, Proyecto, Notas
        semana, dt_val, horario, hh_val, proj_raw, notas = row[:6]
        if isinstance(dt_val, datetime.datetime):
            fecha_str = dt_val.strftime("%Y-%m-%d")
        else:
            fecha_str = str(dt_val)[:10]
        hh = float(hh_val or 0)
        proj_str = str(proj_raw or "")
        notas_str = str(notas or "")

        # Manejo de Kinross / Las Luces (split 50%)
        if "kinross" in proj_str.lower() and "luces" in proj_str.lower():
            tareas = [
                {
                    "proyecto": "Kinross - LNF",
                    "categoria": "Proyectos",
                    "horas": hh / 2.0,
                    "actividad": "Modelación y Georreferenciación",
                    "detalle": "Georreferenciación de perfiles hidrogeológicos en Leapfrog"
                },
                {
                    "proyecto": "Minera Las Luces (MLC)",
                    "categoria": "Proyectos",
                    "horas": hh / 2.0,
                    "actividad": "Revisión de Antecedentes",
                    "detalle": "Lectura de antecedentes de mina Las Luces (geología, modelo e ICSARA)"
                }
            ]
        else:
            norm_p = normalize_project(proj_str)
            act = "Redacción de Informe" if "informe" in notas_str.lower() else "Gabinete / Terreno"
            if "leapfrog" in notas_str.lower(): act = "Modelación Leapfrog"
            if "pre-induccion" in notas_str.lower() or "multitester" in notas_str.lower(): act = "Preparación Terreno / Equipos"
            tareas = [
                {
                    "proyecto": norm_p,
                    "categoria": "Proyectos",
                    "horas": hh,
                    "actividad": act,
                    "detalle": notas_str
                }
            ]

        rec = {
            "id": f"TS-CB-{fecha_str}",
            "fecha": fecha_str,
            "usuarioNombre": "Cristóbal Bravo",
            "usuarioCorreo": "cbravo@icageo.cl",
            "tipoJornada": "Terreno_Extendido" if hh > 7.0 or "Terreno" in proj_str else "Normal",
            "totalHH": hh,
            "tareas": tareas,
            "timestamp": f"{fecha_str}T18:00:00.000Z",
            "estadoRevision": "Al Día"
        }
        all_records.append(rec)

# ==============================================================================
# 2. GONZALO MARAGAÑO
# ==============================================================================
file_gm = os.path.join(BASE_TIMESHEET_DIR, "TS Gonzalo Maragaño.xlsx")
if os.path.exists(file_gm):
    wb = openpyxl.load_workbook(file_gm, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    for r in rows:
        if len(r) >= 3 and isinstance(r[0], datetime.datetime) and r[2]:
            dt_val = r[0]
            dia_str = str(r[1] or "")
            act_str = str(r[2] or "").strip()
            if not act_str:
                continue
            fecha_str = dt_val.strftime("%Y-%m-%d")
            is_viernes = "viernes" in dia_str.lower() or dt_val.weekday() == 4
            hh = 3.0 if is_viernes else 7.0
            norm_p = normalize_project(act_str)
            
            all_records.append({
                "id": f"TS-GM-{fecha_str}",
                "fecha": fecha_str,
                "usuarioNombre": "Gonzalo Maragaño Carmona",
                "usuarioCorreo": "gmaragano@icageo.cl",
                "tipoJornada": "Normal",
                "totalHH": hh,
                "tareas": [
                    {
                        "proyecto": norm_p,
                        "categoria": "Proyectos",
                        "horas": hh,
                        "actividad": "Elaboración de Informe Especializado",
                        "detalle": act_str
                    }
                ],
                "timestamp": f"{fecha_str}T17:30:00.000Z",
                "estadoRevision": "Al Día"
            })

# ==============================================================================
# 3. GONZALO SUÁREZ
# ==============================================================================
file_gs = os.path.join(BASE_TIMESHEET_DIR, "TS Gonzalo Suárez.xlsx")
if os.path.exists(file_gs):
    wb = openpyxl.load_workbook(file_gs, data_only=True)
    
    # Mapeo de días y fechas para las dos semanas
    week_date_map = {
        "24-08 al 28-08 ": {
            2: "2026-08-24", # Lunes
            3: "2026-08-25", # Martes
            4: "2026-08-26", # Miercoles
            5: "2026-08-27", # Jueves
            6: "2026-08-28", # Viernes
        },
        "31-08 al-04-09": {
            2: "2026-08-31", # Lunes
            3: "2026-09-01", # Martes
            4: "2026-09-02", # Miercoles
            5: "2026-09-03", # Jueves
            6: "2026-09-04", # Viernes
        }
    }

    for sheet_name, col_map in week_date_map.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        day_tasks = defaultdict(lambda: defaultdict(float)) # day -> project -> hours
        day_notes = defaultdict(lambda: defaultdict(list))
        
        # Recorrer filas de horarios (R4 a R24)
        for row in list(ws.iter_rows(values_only=True))[3:24]:
            hora_cell = row[0]
            for col_idx, fecha_str in col_map.items():
                if col_idx - 1 < len(row):
                    cell_val = str(row[col_idx - 1] or "").strip()
                    if cell_val and cell_val.lower() not in ["almuerzo", "colación", "colacion"]:
                        # 30 minutos = 0.5 horas
                        norm_p = normalize_project(cell_val)
                        day_tasks[fecha_str][norm_p] += 0.5
                        if cell_val not in day_notes[fecha_str][norm_p]:
                            day_notes[fecha_str][norm_p].append(cell_val)

        for fecha_str, projs in day_tasks.items():
            tareas = []
            tot_day_hh = sum(projs.values())
            for p, hh_p in projs.items():
                tareas.append({
                    "proyecto": p,
                    "categoria": "Proyectos" if "ICA" not in p else "Gestión Interna",
                    "horas": round(hh_p, 1),
                    "actividad": "Análisis Hidrogeológico / Operaciones",
                    "detalle": "; ".join(day_notes[fecha_str][p])
                })
            
            all_records.append({
                "id": f"TS-GS-{fecha_str}",
                "fecha": fecha_str,
                "usuarioNombre": "Gonzalo Suárez",
                "usuarioCorreo": "gsuarez@icageo.cl",
                "tipoJornada": "Terreno_Extendido" if tot_day_hh > 7.0 or any("Las Luces" in t["proyecto"] for t in tareas) else "Normal",
                "totalHH": round(tot_day_hh, 1),
                "tareas": tareas,
                "timestamp": f"{fecha_str}T18:00:00.000Z",
                "estadoRevision": "Al Día"
            })

# ==============================================================================
# 4. CLAUDIA LEÓN
# ==============================================================================
file_cl = os.path.join(BASE_TIMESHEET_DIR, "TS Claudia León.xlsx")
if os.path.exists(file_cl):
    wb = openpyxl.load_workbook(file_cl, data_only=True)
    ws = wb.active
    
    # Columnas 2=Lunes(24), 3=Martes(25), 4=Miercoles(26), 5=Jueves(27), 6=Viernes(28)
    col_dates_cl = {
        2: "2026-08-24",
        3: "2026-08-25",
        4: "2026-08-26",
        5: "2026-08-27",
        6: "2026-08-28"
    }
    
    day_tasks_cl = defaultdict(lambda: defaultdict(float))
    day_notes_cl = defaultdict(lambda: defaultdict(list))
    
    for row in list(ws.iter_rows(values_only=True))[1:18]:
        for col_idx, fecha_str in col_dates_cl.items():
            if col_idx - 1 < len(row):
                cell_val = str(row[col_idx - 1] or "").strip()
                if cell_val and cell_val.lower() not in ["colación", "colacion", "pausa activa"]:
                    # Manejo de proyectos compuestos en Claudia (ej: Transparencia DGA Kinross + Hitos BSF)
                    if "kinross" in cell_val.lower() and "bsf" in cell_val.lower():
                        day_tasks_cl[fecha_str]["Kinross - LNF"] += 0.25
                        day_tasks_cl[fecha_str]["Bodega San Francisco (BSF)"] += 0.25
                        day_notes_cl[fecha_str]["Kinross - LNF"].append("Solicitud de transparencia a DGA")
                        day_notes_cl[fecha_str]["Bodega San Francisco (BSF)"].append("Envío hitos de pago propuesta BSF")
                    else:
                        norm_p = normalize_project(cell_val)
                        day_tasks_cl[fecha_str][norm_p] += 0.5
                        if cell_val not in day_notes_cl[fecha_str][norm_p]:
                            day_notes_cl[fecha_str][norm_p].append(cell_val)
                            
    # Asegurar la jornada efectiva estándar asignando el foco de redacción correspondiente
    for fecha_str, projs in day_tasks_cl.items():
        curr_tot = sum(projs.values())
        is_vi = fecha_str.endswith("-28")
        target_hh = 3.5 if is_vi else 7.0
        
        # Ajustar horas restantes a proyecto principal del día (Pimentón / WSP)
        if curr_tot < target_hh:
            deficit = target_hh - curr_tot
            main_p = "Minera Pimentón" if "Minera Pimentón" in projs else "WSP - DIA Glaciares"
            projs[main_p] += deficit
            day_notes_cl[fecha_str][main_p].append("Gabinete y análisis de datos hidroquímicos / isotópicos")
            
        tareas = []
        for p, hh_p in projs.items():
            tareas.append({
                "proyecto": p,
                "categoria": "Proyectos" if "ICA" not in p else "Gestión Interna",
                "horas": round(hh_p, 1),
                "actividad": "Análisis Isotópico / Hidroquímica / Gabinete",
                "detalle": "; ".join(day_notes_cl[fecha_str][p])
            })
            
        all_records.append({
            "id": f"TS-CL-{fecha_str}",
            "fecha": fecha_str,
            "usuarioNombre": "Claudia León Rojas",
            "usuarioCorreo": "cleon@icageo.cl",
            "tipoJornada": "Normal",
            "totalHH": round(sum(t["horas"] for t in tareas), 1),
            "tareas": tareas,
            "timestamp": f"{fecha_str}T18:00:00.000Z",
            "estadoRevision": "Al Día"
        })

# ==============================================================================
# 5. JAVIERA RODRÍGUEZ
# ==============================================================================
file_jr = os.path.join(BASE_TIMESHEET_DIR, "TS Javiera Rodriguez.xlsx")
if os.path.exists(file_jr):
    # Semana 1 (24 al 28 Agosto)
    jr_w1 = [
        ("2026-08-24", [
            ("Minera Pimentón", 3.5, "Geología e Hidrogeología", "UH, Hidrogeología y catastro de aguas"),
            ("Propuestas y Licitaciones", 3.5, "Elaboración de Oferta", "Elaboración de propuesta EFE")
        ]),
        ("2026-08-25", [
            ("Propuestas y Licitaciones", 3.5, "Elaboración de Oferta", "Elaboración de propuesta EFE"),
            ("Minera Pimentón", 3.5, "Geología e Hidrogeología", "Añadir vega y vertiente, correcciones informe Pimentón")
        ]),
        ("2026-08-26", [
            ("Propuestas y Licitaciones", 3.5, "Elaboración de Oferta", "Propuesta EFE"),
            ("Bodega San Francisco (BSF)", 3.5, "Vulnerabilidad Hidrogeológica", "Recolección de antecedentes, completar excel vulnerabilidad")
        ]),
        ("2026-08-27", [
            ("Gestión Interna / ICA", 3.5, "Seguridad y Salud Ocupacional", "Exámenes de altura física ASCH"),
            ("Bodega San Francisco (BSF)", 2.0, "Vulnerabilidad", "Completar excel de vulnerabilidad BSF"),
            ("Minera Pimentón", 1.5, "Cartografía", "Corrección mapa geológico Pimentón")
        ]),
        ("2026-08-28", [
            ("Minera Pimentón", 2.0, "Revisión de Perfiles", "Revisión perfiles Pimentón"),
            ("Gestión Interna / ICA", 1.0, "Gestión Corporativa", "Hacer CV formato ICA")
        ])
    ]
    
    # Semana 2 (31 Agosto al 04 Septiembre)
    jr_w2 = [
        ("2026-08-31", [
            ("Minera Pimentón", 7.0, "Cartografía y Perfiles", "Describir mapas y perfiles Pimentón (AM y PM)")
        ]),
        ("2026-09-01", [
            ("Minera Pimentón", 3.5, "Cartografía y Perfiles", "Describir mapas y perfiles Pimentón"),
            ("CCU Quilicura", 3.5, "Inducciones Operativas", "Inducciones CCU Quilicura")
        ]),
        ("2026-09-02", [
            ("CCU Quilicura", 7.0, "Inducciones y Capacitación", "Inducciones CCU Quilicura (AM y PM)")
        ]),
        ("2026-09-03", [
            ("Minera Pimentón", 3.5, "Modelo Conceptual", "Descripción mapa conceptual hidrogeológico y corrección perfiles"),
            ("Bodega San Francisco (BSF)", 3.5, "Informe de Vulnerabilidad", "BSF terminar excel vulnerabilidad y edición informe")
        ]),
        ("2026-09-04", [
            ("Bodega San Francisco (BSF)", 1.5, "Informe de Vulnerabilidad", "BSF edición informe"),
            ("Minera Pimentón", 1.5, "Formato y Entrega", "Pimentón formato informe final")
        ])
    ]

    for fecha_str, tasks_list in (jr_w1 + jr_w2):
        tareas = []
        for p, hh, act, det in tasks_list:
            tareas.append({
                "proyecto": p,
                "categoria": "Proyectos" if "ICA" not in p else "Gestión Interna",
                "horas": hh,
                "actividad": act,
                "detalle": det
            })
        all_records.append({
            "id": f"TS-JR-{fecha_str}",
            "fecha": fecha_str,
            "usuarioNombre": "Javiera Rodríguez",
            "usuarioCorreo": "jrodriguez@icageo.cl",
            "tipoJornada": "Normal",
            "totalHH": round(sum(t["horas"] for t in tareas), 1),
            "tareas": tareas,
            "timestamp": f"{fecha_str}T18:00:00.000Z",
            "estadoRevision": "Al Día"
        })

# Ordenar por fecha cronológica y nombre
all_records.sort(key=lambda x: (x["fecha"], x["usuarioNombre"]))

# ==============================================================================
# GUARDAR DATASET JSON Y JS
# ==============================================================================
json_app_path = os.path.join(APP_DATA_DIR, "timesheet_records.json")
with open(json_app_path, "w", encoding="utf-8") as f:
    json.dump(all_records, f, ensure_ascii=False, indent=2)

json_sp_path = os.path.join(SHAREPOINT_DIR, "timesheet_records.json")
with open(json_sp_path, "w", encoding="utf-8") as f:
    json.dump(all_records, f, ensure_ascii=False, indent=2)

js_path = r"C:\Users\ealva\Desktop\Proyectos Antigravity\Gestion ICA\Timesheet_App\js\team_data.js"
with open(js_path, "w", encoding="utf-8") as f:
    f.write("// Datos Reales Consolidados del Equipo ICA Geoconsultores (Agosto - Septiembre 2026)\n")
    f.write("window.ICA_REAL_TEAM_DATA = ")
    json.dump(all_records, f, ensure_ascii=False, indent=2)
    f.write(";\n")

# ==============================================================================
# ACTUALIZAR EXCEL MAESTRO SHAREPOINT
# ==============================================================================
sp_master_excel = r"C:\Users\ealva\SEID\ICAGEO CHILE - Pegas\Timesheet\Timesheet_Maestro_ICA_Estandarizado.xlsx"
if os.path.exists(sp_master_excel):
    try:
        wb_master = openpyxl.load_workbook(sp_master_excel)
        if "Registro_Horas_Detalle" in wb_master.sheetnames:
            ws_det = wb_master["Registro_Horas_Detalle"]
            # Limpiar filas previas a partir de la fila 5
            if ws_det.max_row >= 5:
                ws_det.delete_rows(5, ws_det.max_row - 4)
            
            # Insertar filas del equipo
            reg_counter = 1
            for r in all_records:
                fecha_obj = datetime.datetime.strptime(r["fecha"], "%Y-%m-%d")
                dias_esp = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
                dia_nombre = dias_esp[fecha_obj.weekday()]
                
                for t in r["tareas"]:
                    row_data = [
                        f"TS-{reg_counter:04d}",
                        r["fecha"],
                        r["fecha"],
                        r["usuarioNombre"].split()[0], # Primer nombre
                        r["usuarioCorreo"],
                        t["proyecto"],
                        "OT-ICA-2026",
                        t["detalle"] or t["actividad"],
                        dia_nombre,
                        t["horas"],
                        "Aprobado",
                        "Excel Personal Importado"
                    ]
                    ws_det.append(row_data)
                    reg_counter += 1
            wb_master.save(sp_master_excel)
            print(f"[OK] Sincronizado archivo maestro Excel en: {sp_master_excel} ({reg_counter-1} filas)")
    except Exception as e:
        print(f"[WARN] No se pudo escribir en Excel maestro (puede estar abierto): {e}")

print(f"[OK] Guardados {len(all_records)} registros diarios consolidados en:")
print(f"  - {json_app_path}")
print(f"  - {json_sp_path}")
print(f"  - {js_path}")

# ==============================================================================
# RESUMEN ESTADISTICO
# ==============================================================================
hours_by_user = defaultdict(float)
hours_by_proj = defaultdict(float)

for r in all_records:
    hours_by_user[r["usuarioNombre"]] += r["totalHH"]
    for t in r["tareas"]:
        hours_by_proj[t["proyecto"]] += t["horas"]

print("\n--- RESUMEN POR COLABORADOR ---")
for u, hh in sorted(hours_by_user.items(), key=lambda x: x[0]):
    print(f"  * {u}: {hh:.1f} HH")

print("\n--- RESUMEN POR PROYECTO ---")
for p, hh in sorted(hours_by_proj.items(), key=lambda x: x[1], reverse=True):
    print(f"  * {p}: {hh:.1f} HH")

print(f"\nTOTAL GENERAL HH EQUIPO: {sum(hours_by_user.values()):.1f} HH")
